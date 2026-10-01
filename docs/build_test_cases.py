"""Builds docs/test_cases.xlsx (Phase 1 deliverable).

    python docs/build_test_cases.py --repo https://github.com/<you>/edot-qa-automation

Test data are the exact values used by run C95EA1. Generated data is deterministic per RUN_ID
on the fallback path, so `RUN_ID=C95EA1 python -c "from ai.data_generator import generate_test_data;
print(generate_test_data().data)"` (with no API key) prints these same values.
"""
import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RUN = "C95EA1"
CO = {"name": "PT Aryani Niaga QAC95EA1", "email": "sales.qac95ea1@aryaniniaga.example.com",
      "phone": "88857073301", "street": "Gang Siliwangi No. 135", "industry": "Healthcare"}
CU = {"name": "Toko Sakti Laksita QAC95EA1", "contact": "Ganep Siregar", "phone": "84011073684",
      "email": "owner.qac95ea1@tokoaryaniniaga.example.com", "street": "Jalan Moch. Ramdan No. 91",
      "ktp": "3174010190962441"}
ADDR = "Indonesia / DKI JAKARTA / JAKARTA SELATAN / KEBAYORAN BARU / SENAYAN -> postal code 12190"

COLUMNS = ["Test Case ID", "Title / Description", "Precondition", "Test Steps", "Test Data (exact values)",
           "Expected Result", "Assertion Tier", "Status", "Automated in", "Notes"]

WEB_LOGIN_PRE = "Logged out. eSuite https://esuite.edot.id reachable."
WEB_PRE = "Logged in as the eSuite test account (session shared via storage_state)."
MOB_PRE = "eWork SFA 2.5.0 installed on the emulator; app state cleared (Maestro launchApp clearState)."

CASES = [
    # ---------------- Web ----------------
    ("WEB-LOGIN-01", "Login with valid credentials shows the dashboard", WEB_LOGIN_PRE,
     "1. Open eSuite (redirects to eDOT Account Center)\n2. Click 'Use Email or Username'\n3. Enter email, click Log In\n"
     "4. Enter password, click Log In",
     "Email: value of ESUITE_EMAIL (.env)\nPassword: value of ESUITE_PASSWORD (.env)",
     "Redirected back to https://esuite.edot.id/ and the greeting 'Welcome Back,' is displayed.",
     "Tier 1", "web/tests/test_login.py::test_valid_login_shows_dashboard", "Credentials are never stored in the repo."),
    ("WEB-LOGIN-02", "Login with a wrong password is rejected", WEB_LOGIN_PRE,
     "1. Open eSuite\n2. Click 'Use Email or Username', enter the valid email, Log In\n3. Enter a wrong password, Log In",
     "Email: value of ESUITE_EMAIL\nPassword: WrongPassword123",
     "Error 'Incorrect password' is shown; the dashboard greeting is not shown.",
     "Negative", "web/tests/test_login.py::test_wrong_password_shows_error", ""),
    ("WEB-LOGIN-03", "Login with an unregistered email is stopped", WEB_LOGIN_PRE,
     "1. Open eSuite\n2. Click 'Use Email or Username'\n3. Enter an unregistered email, Log In",
     f"Email: not.registered.{RUN.lower()}@example.com",
     "Dialog 'Email Not Registered' with 'You can continue by creating new account with this email' and the typed "
     "email echoed; no password field.",
     "Negative", "web/tests/test_login.py::test_unregistered_email_shows_prompt",
     "The test clicks nothing in the dialog - no account is created."),
    ("WEB-COMP-01", "Create a new company - it is listed with its ID", WEB_PRE,
     "1. Companies -> '+ Add Company'\n2. Step 1/3: fill every field and the address cascade, Next\n"
     "3. Step 2/3 (legal documents, optional): Next\n4. Step 3/3: Branch Name, 'Fill in with the same data', tick "
     "'Agree to our Policy and Terms', Register\n5. Open Companies",
     f"Company Name: {CO['name']}\nEmail: {CO['email']}\nPhone (+62): {CO['phone']}\nIndustry Type: {CO['industry']}\n"
     f"Company Type: Distributor\nLanguage: English\nStreet Address: {CO['street']}\nAddress: {ADDR}\nBranch Name: Headquarter",
     "Redirected to Companies. Exactly one card with the exact company name, showing the assigned company ID and 'Active'.",
     "Tier 2", "web/tests/test_company.py::test_registered_company_is_listed (registration in the registered_company fixture)",
     "Data is AI-generated (Faker fallback without an API key) and tagged QA<RUN_ID>."),
    ("WEB-COMP-02", "Company detail matches the input field by field", "WEB-COMP-01 passed.",
     "1. Companies -> Manage on the new company\n2. Read every field in Company Details",
     "Same as WEB-COMP-01",
     f"Company Name = {CO['name']}; Company ID = assigned ID; Industry Type = {CO['industry']}; Company Type = Distributor; "
     f"Email = {CO['email']}; Mobile Number = {CO['phone']}; Company Address = {CO['street']}; Province = DKI JAKARTA; "
     "City = JAKARTA SELATAN; District = KEBAYORAN BARU; Sub District = SENAYAN; Postal Code = 12190. "
     "Stored record (company API) has the same values in full (phone as +62...).",
     "Tier 2", "web/tests/test_company.py::test_company_detail_matches_input",
     "The UI truncates dropdown values over 20 chars; the full values are checked via the company API."),
    ("WEB-COMP-03", "Invalid email blocks step 1 of Register Company", WEB_PRE,
     "1. Companies -> '+ Add Company'\n2. Fill step 1 with an invalid email\n3. Click Next",
     f"Email: not-an-email (other fields as WEB-COMP-01)",
     "Message 'Please provide a valid email address'; wizard stays on step 1 ('Register Legal' not shown).",
     "Negative", "web/tests/test_company.py::test_invalid_email_is_rejected", "Nothing is created."),
    ("WEB-COMP-04", "Next stays disabled until step 1 is complete", WEB_PRE,
     "1. Open '+ Add Company'\n2. Check Next\n3. Fill all fields except Street Address, check Next\n4. Fill Street Address",
     f"As WEB-COMP-01; Street Address filled last: {CO['street']}",
     "Next disabled on the empty form and with Street Address empty; enabled once it is filled. Postal Code is derived "
     "(12190) from the Sub District.",
     "Tier 1", "web/tests/test_company.py::test_next_disabled_until_step_is_valid", ""),
    ("WEB-COMP-05", "Company Name is limited to 30 characters", WEB_PRE,
     "1. Open '+ Add Company'\n2. Type a 30-character name\n3. Replace it with a 32-character name",
     "30 chars: PT Tiga Puluh Karakter Pas QA1\n32 chars: PT Tiga Puluh Satu Karakter QA12",
     "30-char name accepted; the 32-char input is refused and the previous value kept.",
     "Negative", "web/tests/test_company.py::test_company_name_is_limited_to_30_characters",
     "No error message exists to assert - product finding BUG-02."),
    ("WEB-COMP-06", "Delete the company - the record is gone", "WEB-COMP-01..02 done (end of session).",
     "1. Manage -> Delete\n2. Tick 'I understand & agree to delete' (Confirm is disabled before)\n3. Confirm",
     "The company created in WEB-COMP-01",
     "Redirected to Companies. GET /api/v1/company/<id> answers 'company not found'; the company ID is no longer in the "
     "account's company list.",
     "Tier 2", "conftest.py::registered_company teardown -> web/company_lifecycle.py::delete_company",
     "The card list keeps showing the deleted company until the next login (BUG-01); the API answers 500 or 404 (BUG-03)."),
    # ---------------- Mobile ----------------
    ("MOB-LOGIN-01", "Valid salesman login shows the dashboard", MOB_PRE,
     "1. Enter Company ID, Username, Password\n2. Tap Sign In",
     "Company ID / Username / Password: EWORK_COMPANY_ID / EWORK_USERNAME / EWORK_PASSWORD from .env "
     "(the fallback account from the brief)",
     "Dashboard shown (main menu grid and 'Today' recap); the Sign In button is gone.",
     "Tier 1", "mobile/flows/login_dashboard.yaml via mobile/tests/test_mobile.py::test_login_shows_dashboard",
     "A company created by the web suite cannot log in to eWork yet - see MOB-LOGIN-02."),
    ("MOB-LOGIN-02", "Handoff: the user of the web-created company is refused", "WEB-COMP-01 created the company; the "
     "suite added a salesman user to it in eSuite (Company User -> + Add User, branch Headquarter).",
     "1. Enter the new company's ID, the new username and its password\n2. Tap Sign In",
     f"Company ID: ID from WEB-COMP-01\nUsername: sales_qa{RUN.lower()}\nPassword: generated per run (git-ignored handoff file)",
     "Dialog 'Oops' with 'Failed to login. Please contact Sales Admin.'; the dashboard is not shown.",
     "Negative", "mobile/flows/login_refused.yaml via test_new_company_user_is_refused",
     "Documents exactly where the web->mobile handoff stops: SFA-side salesman setup is required."),
    ("MOB-CUST-01", "Customer registration cannot be completed without KTP and signature", "MOB-LOGIN-01 passed.",
     "1. New Customer -> New Customer Registration\n2. Complete Basic and Locations, attach a photo\n3. Check Submit\n"
     "4. Enter KTP, Submit\n5. Check Register on the Approval Signature sheet",
     "As MOB-CUST-02; KTP empty, then filled",
     "Submit disabled while KTP is empty, enabled after. Register disabled until the signature is drawn, enabled after.",
     "Tier 1", "mobile/flows/create_customer.yaml (steps before Submit and Register)", ""),
    ("MOB-CUST-02", "Create a customer - it is listed with the entered data", "MOB-LOGIN-01 passed.",
     "1. New Customer -> New Customer Registration\n2. Basic: outlet name, phone, email, contact person, channel, customer "
     "type -> Continue\n3. Locations: address type, address, Province/City/District/Sub district -> Continue\n"
     "4. Documents: take photo, KTP -> Submit\n5. Sign, Register, confirm 'Yes'\n6. 'Data Saved Successfully' -> Continue\n"
     "7. Find the registration in New Customer List",
     f"Outlet Name: {CU['name']}\nPhone (+62): {CU['phone']}\nEmail: {CU['email']}\nContact Person: {CU['contact']}\n"
     f"Channel Type: General Trade (GT)\nCustomer Type: Retailer Small\nAddress Type: Delivery Address\n"
     f"Address: {CU['street']}\nProvince/City/District/Sub district: DKI JAKARTA / JAKARTA SELATAN / KEBAYORAN BARU / "
     f"SENAYAN\nKTP: {CU['ktp']} (synthetic)\nAttachment: photo from the emulator camera",
     f"'Data Saved Successfully'. One card in New Customer List holds together: name {CU['name']}, address {CU['street']}, "
     "group 'Retailer Small', status 'Waiting for Approval', registration number CUST-##### (run C95EA1: CUST-00285).",
     "Tier 2", "mobile/flows/create_customer.yaml via test_create_customer",
     "A pending registration has no detail screen and the app never displays phone/email/contact person, so those three "
     "cannot be asserted in the UI. eWork has no delete for registrations: every one is logged in handoff/created_customers.json."),
]

BUGS = [
    ("BUG-01", "Deleted company stays in the Companies list as 'Active' until the next login",
     "Card list is built from the login token's workspace list; /api/v1/company (header count) is correct."),
    ("BUG-02", "Company Name silently refuses input longer than 30 characters",
     "No maxlength, no message; Next stays disabled with no hint."),
    ("BUG-03", "Lookup of a deleted company answers HTTP 500 'company not found' (sometimes 404)",
     "Seen for companies 5120785 and 5120791."),
]


def build(repo: str, out: Path):
    wb = Workbook()
    header_fill = PatternFill("solid", fgColor="1F4E78")
    wrap = Alignment(wrap_text=True, vertical="top")

    info = wb.active
    info.title = "Info"
    for row in [("eDOT QA Automation take-home - test cases",), ("GitHub repository", repo),
                ("Scope", "Web (eSuite): login, create company, verify detail, delete. Mobile (eWork SFA): login, create customer."),
                ("Test data", f"Exact values used by run {RUN}. Regenerate: RUN_ID={RUN} with no ANTHROPIC_API_KEY."),
                ("Assertion tiers", "Tier 1 = display only. Tier 2 = data created/edited/deleted, values asserted. "
                                    "Negative = the specific error message/state."),
                ("Status column", "Left blank on purpose (filled during execution).")]:
        info.append(row)
    info["A1"].font = Font(bold=True, size=14)
    info.column_dimensions["A"].width = 20
    info.column_dimensions["B"].width = 110
    for r in info.iter_rows(min_row=2):
        r[0].font = Font(bold=True)
        for c in r:
            c.alignment = wrap

    for title, prefix in (("Web", "WEB"), ("Mobile", "MOB")):
        ws = wb.create_sheet(title)
        ws.append(COLUMNS)
        for case in CASES:
            if case[0].startswith(prefix):
                tc_id, name, pre, steps, data, expected, tier, auto, notes = case
                ws.append([tc_id, name, pre, steps, data, expected, tier, "", auto, notes])
        for i, width in enumerate([14, 34, 30, 48, 46, 52, 12, 10, 40, 40], start=1):
            ws.column_dimensions[get_column_letter(i)].width = width
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
        for row in ws.iter_rows(min_row=1):
            for cell in row:
                cell.alignment = wrap
        ws.freeze_panes = "B2"

    bugs = wb.create_sheet("Findings")
    bugs.append(["ID", "Finding", "Evidence / notes"])
    for b in BUGS:
        bugs.append(list(b))
    for i, width in enumerate([10, 70, 80], start=1):
        bugs.column_dimensions[get_column_letter(i)].width = width
    for cell in bugs[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill

    wb.save(out)
    print(f"wrote {out}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--repo", default="https://github.com/<your-username>/edot-qa-automation")
    args = parser.parse_args()
    build(args.repo, Path(__file__).with_name("test_cases.xlsx"))
